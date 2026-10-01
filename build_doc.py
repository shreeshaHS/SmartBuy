import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.parse_xml import parse_xml
import os

doc = docx.Document()

# Page Setup - Normal margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    
    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    hrun1 = hp.add_run("Poornaprajna Institute of Management | Department of Master of Computer Applications\nSmartBuy (Multi-Role E-Commerce Marketplace) - MCA Mini-Project Based Research Monograph")
    hrun1.font.name = "Times New Roman"
    hrun1.font.size = Pt(8.5)
    hrun1.font.italic = True
    hrun1.font.color.rgb = RGBColor(90, 90, 90)

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    frun = fp.add_run("SmartBuy (Bespoke E-Commerce Marketplace Platform) - MCA Mini-Project Based Research Monograph")
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(8)
    frun.font.italic = True
    frun.font.color.rgb = RGBColor(100, 100, 100)

# Colors
MAROON = RGBColor(165, 29, 42)      # #A51D2A
NAVY = RGBColor(24, 43, 73)         # #182B49
DARK_TEXT = RGBColor(33, 37, 41)    # #212529
HEADER_BG = "182B49"                # Navy hex for tables
ALT_ROW_BG = "F4F6F9"               # Light grey-blue for alt rows

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = MAROON
    return p

def add_author_block(name, affiliation, meta):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(name + "\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = DARK_TEXT
    
    r2 = p.add_run(affiliation + "\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(60, 60, 60)
    
    r3 = p.add_run(meta)
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(9)
    r3.font.color.rgb = RGBColor(80, 80, 80)
    p.paragraph_format.space_after = Pt(14)

def add_abstract_heading():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("ABSTRACT")
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = NAVY
    return p

def add_abstract_field(label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    
    r_lbl = p.add_run(label + ": ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(9.5)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = DARK_TEXT
    
    r_txt = p.add_run(text)
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = DARK_TEXT

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = MAROON
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = NAVY
    return p

def add_heading_3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = DARK_TEXT
    return p

def add_body_p(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_TEXT
    return p

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + ": ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = DARK_TEXT
    r_txt = p.add_run(text)
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(10)
    r_txt.font.color.rgb = DARK_TEXT
    return p

def add_table(caption, headers, rows, col_widths=None):
    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(8)
    cp.paragraph_format.space_after = Pt(3)
    cp.paragraph_format.keep_with_next = True
    c_run = cp.add_run(caption)
    c_run.font.name = "Times New Roman"
    c_run.font.size = Pt(9.5)
    c_run.font.bold = True
    c_run.font.color.rgb = NAVY
    
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Format Header Row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_background(cell, HEADER_BG)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    
    # Format Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = ALT_ROW_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row_cells[c_idx]
            if bg_color != "FFFFFF":
                set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(val))
            run.font.name = "Times New Roman"
            run.font.size = Pt(8.5)
            run.font.color.rgb = DARK_TEXT

    # Apply Column Widths if provided
    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Starting document creation...")
